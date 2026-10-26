"""Load the record list from CSV, XLSX, or plain text.

Accepts flexible column names. Minimum: a hostname column. Record type and
target are used when present. Lines/rows that are blank or comments (#) are
skipped.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from .models import HostResult

_HOST_KEYS = {"hostname", "host", "name", "fqdn", "domain", "subdomain", "record"}
_TYPE_KEYS = {"type", "record_type", "rtype", "recordtype"}
_TARGET_KEYS = {"target", "value", "data", "destination", "content", "answer", "ip", "cname"}

_HOST_RE = re.compile(r"^(?=.{1,253}$)(?!-)[A-Za-z0-9_-]{1,63}(?:\.[A-Za-z0-9_-]{1,63})+\.?$")


def _clean_host(h: str) -> str:
    h = h.strip().strip(".").lower()
    # strip scheme / path if someone pasted a URL
    h = re.sub(r"^[a-z]+://", "", h)
    h = h.split("/")[0].split(":")[0]
    return h


def _pick(headers: list[str], keys: set[str]) -> int | None:
    for i, h in enumerate(headers):
        if h.strip().lower() in keys:
            return i
    return None


def load_records(path: str | Path) -> list[HostResult]:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        rows = _read_xlsx(path)
    elif suffix in {".csv", ".tsv"}:
        rows = _read_csv(path, delimiter="\t" if suffix == ".tsv" else ",")
    else:
        rows = _read_txt(path)
    return _rows_to_hosts(rows)


def _read_csv(path: Path, delimiter: str) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [row for row in csv.reader(f, delimiter=delimiter)]


def _read_xlsx(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = []
    for row in ws.iter_rows(values_only=True):
        out.append(["" if c is None else str(c) for c in row])
    wb.close()
    return out


def _read_txt(path: Path) -> list[list[str]]:
    out = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        # allow "host,type,target" or "host type target" or just "host"
        parts = re.split(r"[\s,]+", line)
        out.append(parts)
    return out


def _rows_to_hosts(rows: list[list[str]]) -> list[HostResult]:
    if not rows:
        return []
    # Detect header row
    header = [c.strip().lower() for c in rows[0]]
    has_header = bool(_HOST_KEYS & set(header))
    hi = ti = gi = None
    data = rows
    if has_header:
        hi = _pick(rows[0], _HOST_KEYS)
        ti = _pick(rows[0], _TYPE_KEYS)
        gi = _pick(rows[0], _TARGET_KEYS)
        data = rows[1:]
    else:
        hi = 0
        ti = 1 if any(len(r) > 1 for r in rows) else None
        gi = 2 if any(len(r) > 2 for r in rows) else None

    seen: set[tuple[str, str]] = set()
    hosts: list[HostResult] = []
    for r in data:
        if not r or (hi is not None and hi >= len(r)):
            continue
        raw = r[hi] if hi is not None else r[0]
        if not raw or raw.strip().startswith("#"):
            continue
        host = _clean_host(raw)
        if not host or not _HOST_RE.match(host):
            continue
        rtype = (r[ti].strip().upper() if ti is not None and ti < len(r) else "")
        target = (r[gi].strip() if gi is not None and gi < len(r) else "")
        key = (host, rtype)
        if key in seen:
            continue
        seen.add(key)
        hosts.append(HostResult(hostname=host, record_type=rtype, record_target=target))
    return hosts
