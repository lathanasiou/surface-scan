"""Export results to Excel (multi-sheet) and JSON."""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from .models import HostResult


def export_json(hosts: list[HostResult], path: str | Path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps([h.to_dict() for h in hosts], indent=2))


def _langs(h: HostResult) -> str:
    return ", ".join(f"{d.name}{'/'+d.version if d.version else ''}"
                     for d in h.tech_by_category("language"))


def _cat(h: HostResult, cat: str) -> str:
    return ", ".join(f"{d.name}{'/'+d.version if d.version else ''}"
                     for d in h.tech_by_category(cat))


def _flatten_value(key, value) -> str:
    """Render any HostResult field as a single flat cell, mirroring the JSON."""
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return str(value)
    if key == "detections" and isinstance(value, list):
        parts = []
        for d in value:
            name = d.get("name", "")
            ver = d.get("version")
            cat = d.get("category", "")
            conf = d.get("confidence", "")
            parts.append(f"{name}{'/'+ver if ver else ''} ({cat}, {conf})")
        return " | ".join(parts)
    if key == "advisories" and isinstance(value, dict):
        return " | ".join(f"{aid}: {v.get('verdict','')}" for aid, v in value.items())
    if isinstance(value, (list, tuple)):
        return ", ".join(str(x) for x in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def export_xlsx(hosts: list[HostResult], path: str | Path, advisory_ids: list[str] | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="1F2A44")
    hdr_font = Font(color="FFFFFF", bold=True)

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"

    def autosize(ws, maxw=60):
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, maxw)

    # ---- Summary ----
    ws = wb.active
    ws.title = "Summary"
    live = Counter(h.liveness for h in hosts)
    cdn = Counter(h.cdn_waf for h in hosts)
    lang = Counter()
    for h in hosts:
        for d in h.tech_by_category("language"):
            lang[d.name] += 1
        for d in h.tech_by_category("frontend"):
            lang[f"{d.name} (frontend)"] += 1
    rows = [["Metric", "Value"], ["Total hosts", len(hosts)], ["", ""], ["Liveness", ""]]
    rows += [[f"  {k}", v] for k, v in live.most_common()]
    rows += [["", ""], ["CDN / WAF", ""]]
    rows += [[f"  {k}", v] for k, v in cdn.most_common()]
    rows += [["", ""], ["Technologies", ""]]
    rows += [[f"  {k}", v] for k, v in lang.most_common()]
    for r in rows:
        ws.append(r)
    ws["A1"].font = hdr_font; ws["A1"].fill = hdr_fill
    ws["B1"].font = hdr_font; ws["B1"].fill = hdr_fill
    autosize(ws)

    # ---- All fields (one row per host, every JSON key flattened) ----
    ws = wb.create_sheet("All")
    # Column order = HostResult field order (stable across hosts).
    all_keys: list[str] = []
    for h in hosts:
        for k in h.to_dict().keys():
            if k not in all_keys:
                all_keys.append(k)
    ws.append(all_keys)
    for h in hosts:
        d = h.to_dict()
        ws.append([_flatten_value(k, d.get(k)) for k in all_keys])
    style_header(ws, len(all_keys))
    autosize(ws, maxw=70)

    # ---- Assets ----
    ws = wb.create_sheet("Assets")
    cols = ["hostname", "record_type", "classification", "liveness", "cdn_waf", "cdn_waf_conf",
            "status", "final_url", "title", "resp_ms", "open_ports", "resolved_ips", "asn_org",
            "language", "backend", "frontend", "cms", "server", "tls_issuer", "tls_expiry",
            "dns_notes", "tls_notes"]
    ws.append(cols)
    for h in hosts:
        ws.append([
            h.hostname, h.record_type, h.classification, h.liveness, h.cdn_waf, h.cdn_waf_confidence,
            h.status_code, h.final_url, h.title, h.response_ms,
            ",".join(map(str, h.open_ports)), ",".join(h.resolved_ips), h.asn_org,
            _langs(h), _cat(h, "backend"), _cat(h, "frontend"), _cat(h, "cms"), _cat(h, "server"),
            h.tls_issuer, h.tls_not_after,
            "; ".join(h.dns_notes), "; ".join(h.tls_notes),
        ])
    style_header(ws, len(cols)); autosize(ws)

    # ---- Tech (one row per host+tech) ----
    ws = wb.create_sheet("Tech")
    cols = ["hostname", "category", "technology", "version", "confidence", "evidence"]
    ws.append(cols)
    for h in hosts:
        for d in h.detections:
            ws.append([h.hostname, d.category, d.name, d.version or "", d.confidence,
                       " | ".join(d.evidence)])
    style_header(ws, len(cols)); autosize(ws)

    # ---- DNS ----
    ws = wb.create_sheet("DNS")
    cols = ["hostname", "record_type", "record_target", "cname_chain", "A", "AAAA",
            "MX", "NS", "TXT", "CAA", "notes"]
    ws.append(cols)
    for h in hosts:
        ws.append([h.hostname, h.record_type, h.record_target, " -> ".join(h.cname_chain),
                   ",".join(h.a_records), ",".join(h.aaaa_records), ",".join(h.mx),
                   ",".join(h.ns), " | ".join(h.txt)[:300], ",".join(h.caa),
                   "; ".join(h.dns_notes)])
    style_header(ws, len(cols)); autosize(ws)

    # ---- Flags ----
    ws = wb.create_sheet("Flags")
    cols = ["hostname", "flag", "detail"]
    ws.append(cols)
    warn_fill = PatternFill("solid", fgColor="FFF3CD")
    for h in hosts:
        for n in h.dns_notes:
            ws.append([h.hostname, "dns", n])
        for n in h.tls_notes:
            if "expire" in n.lower() or "expired" in n.lower():
                ws.append([h.hostname, "tls", n])
        if h.cdn_waf == "none" and h.liveness in {"live", "live_error"}:
            ws.append([h.hostname, "no-cdn", "Internet-facing app with no CDN/WAF in front"])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.fill = warn_fill
    style_header(ws, len(cols)); autosize(ws)

    # ---- Advisory sheets ----
    for aid in advisory_ids or []:
        safe = aid.replace("/", "-")[:28]
        ws = wb.create_sheet(f"Adv-{safe}")
        cols = ["hostname", "verdict", "reason", "action"]
        ws.append(cols)
        vcolor = {"affected": "F8D7DA", "likely": "FFE5B4", "possible": "FFF3CD",
                  "unknown": "E2E3E5", "not-detected": "D4EDDA"}
        for h in hosts:
            v = h.advisories.get(aid)
            if not v:
                continue
            ws.append([h.hostname, v["verdict"], v["reason"], v.get("action", "")])
            fill = vcolor.get(v["verdict"])
            if fill:
                for c in ws[ws.max_row]:
                    c.fill = PatternFill("solid", fgColor=fill)
        style_header(ws, len(cols)); autosize(ws)

    wb.save(path)
